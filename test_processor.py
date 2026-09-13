# -*- coding: utf-8 -*-
from __future__ import annotations

import io
import unittest

from openpyxl import Workbook

from processor import (
    extract_asin,
    parse_number,
    process_ads_report,
    process_business_report,
    process_reports,
)


class ExtractAsinTests(unittest.TestCase):
    def test_portfolio_name_keeps_ten_char_asin(self):
        self.assertEqual(extract_asin('B0F5WJBX67 7件套'), 'B0F5WJBX67')

    def test_plain_asin(self):
        self.assertEqual(extract_asin('B07P4LVZNL'), 'B07P4LVZNL')

    def test_missing_asin_is_dropped(self):
        self.assertIsNone(extract_asin('7件套春季促销'))
        self.assertIsNone(extract_asin(''))
        self.assertIsNone(extract_asin(None))


class ParseNumberTests(unittest.TestCase):
    def test_currency_and_percent(self):
        self.assertEqual(parse_number('$1,234.50'), 1234.50)
        self.assertEqual(parse_number('12.5%'), 12.5)
        self.assertEqual(parse_number('(2.00)'), -2.0)
        self.assertEqual(parse_number('—'), 0.0)
        self.assertEqual(parse_number('1.234,56'), 1234.56)


class BusinessReportTests(unittest.TestCase):
    def test_filter_zero_and_merge_asin(self):
        csv_text = (
            '(子)ASIN,标题,会话数 - 总计,转化率 - 总计,页面浏览量 - 总计,'
            '页面浏览量百分比 - 总计,推荐报价（推荐报价展示位）百分比,已订购商品数量,'
            '商品会话百分比,已订购商品销售额,订单商品总数\n'
            'B0F5WJBX67,A,10,1,1,1,1,2,20%,10.5,1\n'
            'B0F5WJBX67,A2,5,1,1,1,1,3,60%,4.5,1\n'
            'B0AAA00001,B,9,1,1,1,1,0,0%,0,0\n'
            'B0BBB00002,C,8,1,1,1,1,4,50%,20,1\n'
        )
        frame = process_business_report(csv_text.encode('utf-8-sig'), 'biz.csv')
        self.assertEqual(list(frame.columns), [
            '(子)ASIN', '会话数 - 总计', '已订购商品数量', '商品会话百分比', '已订购商品销售额',
        ])
        self.assertEqual(frame['(子)ASIN'].tolist(), ['B0F5WJBX67', 'B0BBB00002'])
        self.assertEqual(frame.loc[0, '会话数 - 总计'], 15)
        self.assertEqual(frame.loc[0, '已订购商品数量'], 5)
        self.assertEqual(frame.loc[0, '已订购商品销售额'], 15.0)
        self.assertEqual(frame.loc[0, '商品会话百分比'], 0.3333)


class AdsReportTests(unittest.TestCase):
    def test_extract_asin_filter_zero_and_merge(self):
        csv_text = (
            ',广告组合,状态,预算类型,预算(USD),预算开始日期,预算结束日期,广告活动数量,'
            '展示次数,点击量,CTR,支出(USD),CPC(USD),订单数量,销量(USD),ACOS,'
            '广告投资回报率 (ROAS),NTB订单数量,NTB 订单数量百分比,NTB 销售额(USD),'
            '品牌新客销售额比例,可见展示量,VCPM(USD)\n'
            ',B0F5WJBX67 7件套,enabled,x,1,1,1,1,1,1,1,2.5,1,1,1,1,1,1,1,1,1,1,1\n'
            ',B0F5WJBX67 备用,enabled,x,1,1,1,1,1,1,1,1.5,1,1,1,1,1,1,1,1,1,1,1\n'
            ',春季促销无编码,enabled,x,1,1,1,1,1,1,1,9,1,1,1,1,1,1,1,1,1,1,1\n'
            ',B0CCC00003,enabled,x,1,1,1,1,1,1,1,0,1,1,1,1,1,1,1,1,1,1,1\n'
            ',B0DDD00004 套装,enabled,x,1,1,1,1,1,1,1,4,1,1,1,1,1,1,1,1,1,1,1\n'
        )
        frame = process_ads_report(csv_text.encode('utf-8-sig'), 'ads.csv')
        self.assertEqual(list(frame.columns), ['广告组合', '支出(USD)'])
        self.assertEqual(frame['广告组合'].tolist(), ['B0F5WJBX67', 'B0DDD00004'])
        self.assertEqual(frame.loc[0, '支出(USD)'], 4.0)


class TemplateWriteTests(unittest.TestCase):
    def _sample_files(self):
        business_csv = (
            '(子)ASIN,标题,会话数 - 总计,转化率 - 总计,页面浏览量 - 总计,'
            '页面浏览量百分比 - 总计,推荐报价（推荐报价展示位）百分比,已订购商品数量,'
            '商品会话百分比,已订购商品销售额,订单商品总数\n'
            'B0F5WJBX67,A,10,1,1,1,1,2,20%,10,1\n'
            'B0ZZZ99999,Z,8,1,1,1,1,1,12%,3,1\n'
        ).encode('utf-8-sig')
        ads_csv = (
            '广告组合,支出(USD)\n'
            'B0F5WJBX67 7件套,2.2\n'
            'B0EEE00005 套装,3.3\n'
        ).encode('utf-8-sig')
        return business_csv, ads_csv

    def test_write_from_j2_o2_and_highlight_unknown_asin(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'US'
        sheet['B2'] = 'ASIN'
        sheet['B3'] = 'B0F5WJBX67'
        sheet['B4'] = 'B0BBB00002'
        buffer = io.BytesIO()
        workbook.save(buffer)
        business_csv, ads_csv = self._sample_files()

        output, summary = process_reports(
            business_csv, 'biz.csv',
            ads_csv, 'ads.csv',
            buffer.getvalue(), 'template.xlsx',
        )
        from openpyxl import load_workbook
        result = load_workbook(io.BytesIO(output))
        ws = result['US']
        self.assertEqual(ws['J2'].value, 'B0F5WJBX67')
        self.assertEqual(ws['J3'].value, 'B0ZZZ99999')
        self.assertEqual(ws['O2'].value, 'B0F5WJBX67')
        self.assertEqual(ws['O3'].value, 'B0EEE00005')
        self.assertEqual(ws['J2'].fill.fgColor.rgb, '00000000')
        self.assertEqual(ws['J3'].fill.fgColor.rgb, '00FFFF00')
        self.assertEqual(ws['O3'].fill.fgColor.rgb, '00FFFF00')
        self.assertEqual(summary['highlighted_j'], ['B0ZZZ99999'])
        self.assertEqual(summary['highlighted_o'], ['B0EEE00005'])

    def test_only_us_sheet_is_modified(self):
        workbook = Workbook()
        refund = workbook.active
        refund.title = '回款'
        refund['B3'] = 'should-not-change'
        refund['J2'] = 'refund-j2'
        us = workbook.create_sheet('US')
        us['B2'] = 'ASIN'
        us['B3'] = 'B0F5WJBX67'
        us['J1'] = '（子）ASIN'
        us['O1'] = '广告组合'
        uk = workbook.create_sheet('UK')
        uk['B3'] = 'B0UKUKUKUK'
        uk['J2'] = 'uk-keep'
        uk['O2'] = 'uk-ads'
        buffer = io.BytesIO()
        workbook.save(buffer)
        business_csv, ads_csv = self._sample_files()

        output, _summary = process_reports(
            business_csv, 'biz.csv',
            ads_csv, 'ads.csv',
            buffer.getvalue(), 'template.xlsx',
        )
        from openpyxl import load_workbook
        result = load_workbook(io.BytesIO(output))
        self.assertEqual(result.sheetnames, ['回款', 'US', 'UK'])
        self.assertEqual(result['回款']['J2'].value, 'refund-j2')
        self.assertEqual(result['回款']['B3'].value, 'should-not-change')
        self.assertEqual(result['UK']['J2'].value, 'uk-keep')
        self.assertEqual(result['UK']['O2'].value, 'uk-ads')
        self.assertEqual(result['US']['J1'].value, '（子）ASIN')
        self.assertEqual(result['US']['J2'].value, 'B0F5WJBX67')
        self.assertEqual(result['US']['O2'].value, 'B0F5WJBX67')

    def test_missing_us_sheet_raises(self):
        workbook = Workbook()
        workbook.active.title = '回款'
        buffer = io.BytesIO()
        workbook.save(buffer)
        business_csv, ads_csv = self._sample_files()
        with self.assertRaisesRegex(ValueError, 'US'):
            process_reports(
                business_csv, 'biz.csv',
                ads_csv, 'ads.csv',
                buffer.getvalue(), 'template.xlsx',
            )


if __name__ == '__main__':
    unittest.main()
