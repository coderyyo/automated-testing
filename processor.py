# -*- coding: utf-8 -*-
"""亚马逊业务报表 / 广告报表处理，并写入模板 Excel。"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

# 商品 ASIN：固定 10 位，主流为 B + 9 位字母数字，例如 B0F5WJBX67
ASIN_B_PATTERN = re.compile(r'(?<![A-Z0-9])(B[0-9A-Z]{9})(?![A-Z0-9])', re.IGNORECASE)
ASIN_GENERIC_PATTERN = re.compile(r'(?<![A-Z0-9])([A-Z0-9]{10})(?![A-Z0-9])', re.IGNORECASE)

HIGHLIGHT_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
EMPTY_FILL = PatternFill(fill_type=None)

BUSINESS_COLUMNS = [
    '(子)ASIN',
    '会话数 - 总计',
    '已订购商品数量',
    '商品会话百分比',
    '已订购商品销售额',
]
ADS_COLUMNS = [
    '广告组合',
    '支出(USD)',
]

BUSINESS_ALIASES = {
    '(子)ASIN': ['(子)asin', '子asin', 'asin', 'childasin', '子asin(子)'],
    '会话数 - 总计': ['会话数-总计', '会话数总计', 'sessions-total', 'sessionstotal'],
    '已订购商品数量': ['已订购商品数量', 'unitsordered', '已订购商品数'],
    '商品会话百分比': ['商品会话百分比', 'unitsessionpercentage', 'unit会话百分比'],
    '已订购商品销售额': ['已订购商品销售额', 'orderedproductsales', '已订购商品销售'],
}
ADS_ALIASES = {
    '广告组合': ['广告组合', 'portfolio', 'portfolioname'],
    '支出(USD)': ['支出(usd)', 'spend(usd)', '支出usd', 'spend'],
}


def normalize_header(name: Any) -> str:
    text = str(name).strip().lower()
    text = text.replace('（', '(').replace('）', ')')
    text = text.replace('_x000d_', '').replace('\n', '').replace('\r', '')
    text = re.sub(r'\s+', '', text)
    if text.startswith('unnamed:'):
        return ''
    return text


def extract_asin(value: Any) -> Optional[str]:
    """从单元格稳定提取 10 位 ASIN。优先 B 开头的商品 ASIN，抽不到则返回 None。"""
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    text = str(value).strip().upper()
    if text in ('', 'NAN', 'NONE', 'NULL', '-'):
        return None

    exact_b = re.fullmatch(r'B[0-9A-Z]{9}', text)
    if exact_b:
        return exact_b.group(0)

    match_b = ASIN_B_PATTERN.search(text)
    if match_b:
        return match_b.group(1).upper()

    first_token = re.split(r'[\s,，/|]+', text, maxsplit=1)[0]
    if re.fullmatch(r'B[0-9A-Z]{9}', first_token):
        return first_token

    match_generic = ASIN_GENERIC_PATTERN.search(text)
    if match_generic:
        candidate = match_generic.group(1).upper()
        has_letter = re.search(r'[A-Z]', candidate) is not None
        has_digit = re.search(r'[0-9]', candidate) is not None
        if has_letter and has_digit:
            return candidate
    return None


def parse_number(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        if pd.isna(value):
            return 0.0
        return float(value)

    text = str(value).strip()
    if text in ('', '-', '--', '—', 'N/A', 'NA', 'nan', 'None', 'null'):
        return 0.0

    negative = text.startswith('(') and text.endswith(')')
    if negative:
        text = text[1:-1]

    text = re.sub(r'(?i)usd|gbp|eur|us\$|\$|￥|¥|€|£', '', text)
    text = text.replace('%', '')
    text = text.strip()
    text = re.sub(r'[^\d.,\-]', '', text)
    if text in ('', '-', '.', ',', '-.', '-,', '.-'):
        return 0.0
    if ',' in text and '.' in text:
        if text.rfind(',') > text.rfind('.'):
            text = text.replace('.', '').replace(',', '.')
        else:
            text = text.replace(',', '')
    elif ',' in text:
        parts = text.split(',')
        if len(parts) == 2 and len(parts[-1]) in (1, 2):
            text = text.replace(',', '.')
        else:
            text = text.replace(',', '')
    elif text.count('.') > 1:
        text = text.replace('.', '')
    try:
        number = float(text)
    except ValueError:
        return 0.0
    if negative:
        number = -number
    return number


def is_zero(value: float) -> bool:
    return abs(value) < 1e-9


def find_column(columns: Sequence[Any], aliases: Sequence[str]) -> Optional[Any]:
    normalized_aliases = [normalize_header(item) for item in aliases]
    mapping = {normalize_header(col): col for col in columns if normalize_header(col)}
    for alias in normalized_aliases:
        if alias in mapping:
            return mapping[alias]
    for alias in normalized_aliases:
        for norm_name, original in mapping.items():
            if alias and alias in norm_name:
                return original
    return None


def _decode_bytes(raw: bytes) -> str:
    if raw.startswith(b'\xff\xfe') or raw.startswith(b'\xfe\xff'):
        return raw.decode('utf-16')
    for encoding in ('utf-8-sig', 'utf-8', 'gb18030', 'utf-16', 'utf-16-le', 'latin1'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', errors='replace')


def _header_row_index(lines: List[str], hints: Sequence[str]) -> int:
    normalized_hints = [normalize_header(hint) for hint in hints]
    for index, line in enumerate(lines[:40]):
        compact = normalize_header(line.replace(',', ' ').replace('\t', ' ').replace(';', ' '))
        for hint in normalized_hints:
            if hint and hint in compact:
                return index
    return 0


def _read_csv_text(text: str, hints: Sequence[str]) -> pd.DataFrame:
    lines = text.splitlines()
    if not lines:
        return pd.DataFrame()
    start = _header_row_index(lines, hints)
    sliced = '\n'.join(lines[start:])
    buffer = io.StringIO(sliced)
    try:
        frame = pd.read_csv(buffer, sep=None, engine='python')
    except Exception:
        buffer = io.StringIO(sliced)
        frame = pd.read_csv(buffer)
    if frame.shape[1] == 1:
        buffer = io.StringIO(sliced)
        frame = pd.read_csv(buffer, sep='\t')
    return frame


def load_table(raw: bytes, filename: str, hints: Sequence[str]) -> pd.DataFrame:
    name = (filename or '').lower()
    if name.endswith(('.xlsx', '.xlsm', '.xls')):
        frame = pd.read_excel(io.BytesIO(raw))
        if frame.empty:
            return frame
        compact_headers = ''.join(normalize_header(col) for col in frame.columns)
        needed = [normalize_header(hint) for hint in hints]
        if any(item and item in compact_headers for item in needed):
            return frame
        raw_frame = pd.read_excel(io.BytesIO(raw), header=None)
        for index in range(min(40, len(raw_frame))):
            row_text = ' '.join(str(item) for item in raw_frame.iloc[index].tolist())
            compact = normalize_header(row_text)
            if any(normalize_header(hint) in compact for hint in hints if hint):
                header = raw_frame.iloc[index].tolist()
                body = raw_frame.iloc[index + 1:].copy()
                body.columns = header
                return body.reset_index(drop=True)
        return frame
    text = _decode_bytes(raw)
    return _read_csv_text(text, hints)


def _require_columns(frame: pd.DataFrame, required: Dict[str, Sequence[str]], report_name: str) -> Dict[str, Any]:
    found = {}
    missing = []
    for logical_name, aliases in required.items():
        column = find_column(frame.columns, aliases)
        if column is None:
            missing.append(logical_name)
        else:
            found[logical_name] = column
    if missing:
        available = [str(col) for col in frame.columns if str(col).strip() and not str(col).lower().startswith('unnamed')]
        raise ValueError(
            '{0}缺少列：{1}。当前识别到的表头：{2}'.format(
                report_name,
                '、'.join(missing),
                '、'.join(available) if available else '无',
            )
        )
    return found


def process_business_report(raw: bytes, filename: str) -> pd.DataFrame:
    frame = load_table(raw, filename, BUSINESS_COLUMNS)
    columns = _require_columns(frame, BUSINESS_ALIASES, '业务报表')

    work = pd.DataFrame({
        'asin_raw': frame[columns['(子)ASIN']],
        'sessions': frame[columns['会话数 - 总计']].map(parse_number),
        'units': frame[columns['已订购商品数量']].map(parse_number),
        'sales': frame[columns['已订购商品销售额']].map(parse_number),
    })
    work['asin'] = work['asin_raw'].map(extract_asin)
    work = work[work['asin'].notna()].copy()
    work = work[~work['units'].map(is_zero)].copy()
    if work.empty:
        return pd.DataFrame(columns=BUSINESS_COLUMNS)

    work['_order'] = range(len(work))
    grouped = work.groupby('asin', as_index=False).agg({
        'sessions': 'sum',
        'units': 'sum',
        'sales': 'sum',
        '_order': 'min',
    }).sort_values('_order', kind='mergesort')

    def unit_session_pct(row: pd.Series) -> float:
        if is_zero(row['sessions']):
            return 0.0
        return round(row['units'] / row['sessions'], 4)

    grouped['pct'] = grouped.apply(unit_session_pct, axis=1)
    result = pd.DataFrame({
        '(子)ASIN': grouped['asin'],
        '会话数 - 总计': grouped['sessions'],
        '已订购商品数量': grouped['units'],
        '商品会话百分比': grouped['pct'],
        '已订购商品销售额': grouped['sales'].map(lambda value: round(float(value), 2)),
    })
    return result.reset_index(drop=True)


def process_ads_report(raw: bytes, filename: str) -> pd.DataFrame:
    frame = load_table(raw, filename, ADS_COLUMNS)
    columns = _require_columns(frame, ADS_ALIASES, '广告报表')

    work = pd.DataFrame({
        'portfolio': frame[columns['广告组合']],
        'spend': frame[columns['支出(USD)']].map(parse_number),
    })
    work['asin'] = work['portfolio'].map(extract_asin)
    work = work[work['asin'].notna()].copy()
    work = work[~work['spend'].map(is_zero)].copy()
    if work.empty:
        return pd.DataFrame(columns=ADS_COLUMNS)

    work['_order'] = range(len(work))
    grouped = work.groupby('asin', as_index=False).agg({
        'spend': 'sum',
        '_order': 'min',
    }).sort_values('_order', kind='mergesort')

    result = pd.DataFrame({
        '广告组合': grouped['asin'],
        '支出(USD)': grouped['spend'].map(lambda value: round(float(value), 2)),
    })
    return result.reset_index(drop=True)


def _clean_number_for_excel(value: Any) -> Any:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def get_us_worksheet(workbook) -> Worksheet:
    for name in workbook.sheetnames:
        if str(name).strip().upper() == 'US':
            return workbook[name]
    raise ValueError(
        '模板中没有找到名为 US 的工作表，当前工作表：{0}'.format(
            '、'.join(workbook.sheetnames) if workbook.sheetnames else '无'
        )
    )


def collect_template_asins(worksheet: Worksheet) -> set:
    asins = set()
    for row in range(3, worksheet.max_row + 1):
        asin = extract_asin(worksheet.cell(row=row, column=2).value)
        if asin is not None:
            asins.add(asin)
    return asins


def _clear_block(worksheet: Worksheet, min_col: int, max_col: int, start_row: int) -> None:
    last_row = worksheet.max_row if worksheet.max_row >= start_row else start_row
    for row in range(start_row, last_row + 1):
        has_value = False
        for col in range(min_col, max_col + 1):
            if worksheet.cell(row=row, column=col).value not in (None, ''):
                has_value = True
                break
        if not has_value:
            continue
        for col in range(min_col, max_col + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.value = None
            if col in (10, 15):
                cell.fill = EMPTY_FILL


def _write_frame(
    worksheet: Worksheet,
    frame: pd.DataFrame,
    start_row: int,
    start_col: int,
    template_asins: set,
    asin_col_offset: int,
) -> List[str]:
    highlighted = []
    for index, row in enumerate(frame.itertuples(index=False, name=None)):
        excel_row = start_row + index
        for offset, value in enumerate(row):
            cell = worksheet.cell(row=excel_row, column=start_col + offset)
            cell.value = _clean_number_for_excel(value)
            if offset == asin_col_offset:
                asin = extract_asin(value)
                if asin is not None and asin not in template_asins:
                    cell.fill = HIGHLIGHT_FILL
                    highlighted.append(asin)
                else:
                    cell.fill = EMPTY_FILL
    return highlighted


def write_to_template(template_raw: bytes, business_df: pd.DataFrame, ads_df: pd.DataFrame) -> Tuple[bytes, Dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(template_raw))
    worksheet = get_us_worksheet(workbook)
    template_asins = collect_template_asins(worksheet)

    _clear_block(worksheet, 10, 14, 2)
    _clear_block(worksheet, 15, 16, 2)

    highlighted_j = _write_frame(worksheet, business_df, 2, 10, template_asins, 0)
    highlighted_o = _write_frame(worksheet, ads_df, 2, 15, template_asins, 0)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    summary = {
        'business_rows': int(len(business_df)),
        'ads_rows': int(len(ads_df)),
        'template_asins': int(len(template_asins)),
        'highlighted_j': highlighted_j,
        'highlighted_o': highlighted_o,
    }
    return output.getvalue(), summary


def process_reports(
    business_raw: bytes,
    business_name: str,
    ads_raw: bytes,
    ads_name: str,
    template_raw: bytes,
    template_name: str,
) -> Tuple[bytes, Dict[str, Any]]:
    if not template_name.lower().endswith(('.xlsx', '.xlsm')):
        raise ValueError('模板表请上传 .xlsx 或 .xlsm 文件')
    business_df = process_business_report(business_raw, business_name)
    ads_df = process_ads_report(ads_raw, ads_name)
    return write_to_template(template_raw, business_df, ads_df)
